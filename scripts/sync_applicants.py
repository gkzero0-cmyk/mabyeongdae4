from __future__ import annotations

import json
import os
import re
import time
import urllib.request
from copy import copy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / 'mabyeongdae4_applicant_classification_2026-09-09_v3.xlsx'
TYPE_MAP_PATH = ROOT / 'applicant-types-v3.js'
SITE_API = os.environ.get('MABYEONGDAE4_COMMENTS_API', 'https://mabyeongdae4.vercel.app/api/comments')
POST_URL = 'https://www.sooplive.com/station/devil0108/post/206507027'

VALID_KR_TYPES = {'병사', '간부', '미분류'}
TYPE_TO_JS = {'병사': 'soldier', '간부': 'officer', '미분류': 'unknown'}

# User-confirmed manual decisions. These always outrank automatic detection.
MANUAL_FINAL_OVERRIDES = {
    'gofl2237': '미분류',
    'jaeparkk': '간부',
    'heda221112': '병사',
    'kimtalggy': '병사',
    'lsh8071': '병사',
    'toocat030': '미분류',
    'dnwnwjdqhr53': '미분류',
    'niniming': '미분류',
    'oiguu5252': '미분류',
    'kimhaetae': '미분류',
    'chiya1207': '미분류',
    'doramzi610': '병사',
}


def _soldier_role(text: str) -> bool:
    value = str(text or '')
    return bool(re.search(r'훈련병|훈병|행정병|병사', value, re.I) or re.search(r'(?:^|[^\w가-힣])병(?=$|[^\w가-힣])', value))


def _strip_negative_officer(text: str) -> str:
    return re.sub(
        r'간부\s*(?:는\s*)?(?:없이|[xX]|아님|말고|제외|지원\s*안\s*함|신청\s*안\s*함|안\s*(?:함|합니다|해요|할게요|하겠습니다))',
        ' ',
        str(text or ''),
        flags=re.I,
    )


def _mixed_roles(text: str) -> bool:
    value = str(text or '')
    return bool(
        re.search(r'간부\s*(?:OR|또는|혹은|&|/|\|)\s*(?:행정병|훈련병|훈병|병사|병)', value, re.I)
        or re.search(r'(?:행정병|훈련병|훈병|병사|병)\s*(?:OR|또는|혹은|&|/|\|)\s*간부', value, re.I)
    )


def classify_comment(comment: str) -> str:
    raw = str(comment or '').strip()
    if not raw:
        return '미분류'

    # The first explicit role token wins. This intentionally handles
    # "-병사 -후추 하겠습니다~!" as 병사 instead of treating every 후추 comment as unknown.
    first_line = next((line.strip() for line in raw.splitlines() if line.strip()), '')
    first_line = re.sub(r'^[^\w가-힣]*', '', _strip_negative_officer(first_line)).strip()
    if first_line:
        if _mixed_roles(first_line):
            return '미분류'
        match = re.match(r'^(?:일반\s*)?(간부|행정병|훈련병|훈병|병사|병)(?=$|[\s/|,;:：()\[\]{}<>·&-])', first_line, re.I)
        if match:
            return '간부' if match.group(1) == '간부' else '병사'

    text = re.sub(r'\s+', ' ', _strip_negative_officer(raw)).strip()

    field_types = set()
    for match in re.finditer(r'(?:신청|지원)\s*분야\s*(?:[:：/|\-]\s*)?([^\n]{1,80})', text, re.I):
        value = match.group(1).strip()
        value = re.split(r'(?:(?:마크\s*서버|마크서버|마병대)\s*경험|자기\s*소개|지원\s*동기|특이사항)\s*[:：]', value, maxsplit=1, flags=re.I)[0]
        has_officer = bool(re.search(r'간부', value, re.I))
        has_soldier = _soldier_role(value)
        if has_officer and has_soldier:
            return '미분류'
        if has_officer:
            field_types.add('간부')
        if has_soldier:
            field_types.add('병사')
    if len(field_types) > 1:
        return '미분류'
    if len(field_types) == 1:
        return next(iter(field_types))

    has_officer_intent = bool(re.search(r'간부\s*(?:로\s*)?(?:신청|지원)', text, re.I))
    has_soldier_intent = bool(
        re.search(r'(?:훈련병|훈병|행정병|병사)\s*(?:로\s*)?(?:신청|지원)', text, re.I)
        or re.search(r'(?:^|[^\w가-힣])병\s*(?:(?:으로|로)\s*)?(?:신청|지원)', text, re.I)
    )
    if has_officer_intent and has_soldier_intent:
        return '미분류'
    if has_officer_intent:
        return '간부'
    if has_soldier_intent:
        return '병사'

    if _mixed_roles(text):
        return '미분류'

    return '미분류'


def choose_final_type(auto_type: str, existing_final: str = '') -> str:
    existing = str(existing_final or '').strip()
    if existing in VALID_KR_TYPES:
        return existing
    return auto_type if auto_type in VALID_KR_TYPES else '미분류'


def _comment_sort_key(item: dict) -> tuple[str, int]:
    reg_date = str(item.get('regDate') or '')
    try:
        comment_no = int(str(item.get('commentNo') or '0'))
    except ValueError:
        comment_no = 0
    return (reg_date, comment_no)


def dedupe_comments_by_user(comments: list[dict]) -> list[dict]:
    best: dict[str, dict] = {}
    order: list[str] = []
    for item in comments or []:
        user_id = str(item.get('userId') or '').strip().lower()
        if not user_id:
            continue
        if user_id not in best:
            best[user_id] = item
            order.append(user_id)
        elif _comment_sort_key(item) > _comment_sort_key(best[user_id]):
            best[user_id] = item
    return [best[user_id] for user_id in order]


def fetch_comments(url: str = SITE_API, attempts: int = 5) -> list[dict]:
    last_error = None
    for attempt in range(attempts):
        try:
            request = urllib.request.Request(url, headers={'User-Agent': 'mabyeongdae4-xlsx-sync/1.0', 'Accept': 'application/json'})
            with urllib.request.urlopen(request, timeout=45) as response:
                payload = json.load(response)
            if not payload.get('ok') or not isinstance(payload.get('comments'), list):
                raise RuntimeError(payload.get('error') or 'invalid comments payload')
            return dedupe_comments_by_user(payload['comments'])
        except Exception as exc:  # noqa: BLE001 - retry network/provider errors
            last_error = exc
            if attempt + 1 < attempts:
                time.sleep(min(2 ** attempt, 8))
    raise RuntimeError(f'failed to fetch comments after {attempts} attempts: {last_error}')


def build_comment_url(item: dict) -> str:
    explicit = str(item.get('commentUrl') or '').strip()
    if explicit:
        return explicit
    comment_no = str(item.get('commentNo') or '').strip()
    return f'{POST_URL}#comment_noti{comment_no}' if comment_no else POST_URL


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int = 9) -> None:
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _existing_rows(ws) -> tuple[list[str], dict[str, dict]]:
    order = []
    rows = {}
    for row in range(2, ws.max_row + 1):
        user_id = str(ws.cell(row, 3).value or '').strip().lower()
        if not user_id:
            continue
        order.append(user_id)
        rows[user_id] = {
            'row': row,
            'nick': str(ws.cell(row, 2).value or ''),
            'auto': str(ws.cell(row, 4).value or ''),
            'final': str(ws.cell(row, 5).value or ''),
            'status': str(ws.cell(row, 6).value or ''),
            'note': str(ws.cell(row, 7).value or ''),
            'station': str(ws.cell(row, 8).value or ''),
            'comment_url': str(ws.cell(row, 9).value or ''),
        }
    return order, rows


def _target_rows(comments: list[dict], existing_order: list[str], existing: dict[str, dict]) -> list[list]:
    live = {str(item.get('userId') or '').strip().lower(): item for item in comments}
    ids = [user_id for user_id in existing_order if user_id in live]
    ids.extend(user_id for user_id in live if user_id not in existing)

    out = []
    for number, user_id in enumerate(ids, start=1):
        item = live[user_id]
        auto_type = classify_comment(str(item.get('comment') or ''))
        current = existing.get(user_id, {})
        manual = MANUAL_FINAL_OVERRIDES.get(user_id)
        if manual:
            final_type = manual
        else:
            # The workbook is authoritative for applicants already reviewed.
            # Re-evaluate only the automatic column while preserving the existing final decision.
            final_type = choose_final_type(auto_type, current.get('final', ''))
        status = '일치' if final_type == auto_type else '수정 필요'
        note = current.get('note', '') if status == '수정 필요' else ''
        if user_id == 'doramzi610' and status == '수정 필요':
            note = "신청댓글에 '-병사' 명시 → 병사"
        out.append([
            number,
            str(item.get('userNick') or user_id).strip(),
            user_id,
            auto_type,
            final_type,
            status,
            note,
            f'https://www.sooplive.com/station/{user_id}',
            build_comment_url(item),
        ])
    return out


def render_type_map(rows: list[list]) -> str:
    entries = []
    for row in rows:
        user_id = str(row[2]).strip().lower()
        js_type = TYPE_TO_JS.get(str(row[4]).strip(), 'unknown')
        entries.append(f'  {json.dumps(user_id, ensure_ascii=False)}: {json.dumps(js_type)}')
    body = ',\n'.join(entries)
    count = len(entries)
    return (
        "(function (root, factory) {\n"
        "  const map = factory();\n"
        "  if (typeof module === 'object' && module.exports) module.exports = map;\n"
        "  if (root) root.Mabyeongdae4ApplicantTypesV3 = map;\n"
        "})(typeof globalThis !== 'undefined' ? globalThis : this, function () {\n"
        "  // Source: mabyeongdae4_applicant_classification_2026-09-09_v3.xlsx\n"
        f"  // {count} unique SOOP IDs. This file is the authoritative 병사/간부/미분류 source.\n"
        "  return Object.freeze({\n"
        f"{body}\n"
        "});\n"
        "});\n"
    )


def sync_workbook(comments: list[dict] | None = None, workbook_path: Path = WORKBOOK_PATH, type_map_path: Path = TYPE_MAP_PATH) -> bool:
    # Imported lazily so the pure classification tests do not depend on openpyxl.
    from openpyxl import load_workbook

    comments = comments if comments is not None else fetch_comments()
    wb = load_workbook(workbook_path)
    ws = wb['신청자 전체']
    summary = wb['분류 요약']
    existing_order, existing = _existing_rows(ws)
    rows = _target_rows(comments, existing_order, existing)

    current_rows = []
    for row in range(2, ws.max_row + 1):
        if not ws.cell(row, 3).value:
            continue
        current_rows.append([ws.cell(row, col).value if ws.cell(row, col).value is not None else '' for col in range(1, 10)])

    type_map_text = render_type_map(rows)
    existing_map_text = type_map_path.read_text(encoding='utf-8') if type_map_path.exists() else ''
    changed = current_rows != rows or existing_map_text != type_map_text
    if not changed:
        return False

    start_row = 2
    old_max = max(ws.max_row, start_row + len(rows) - 1)
    if rows:
        for target_row in range(start_row, start_row + len(rows)):
            if target_row > ws.max_row:
                _copy_row_style(ws, max(2, target_row - 1), target_row)
        for r_index, row_values in enumerate(rows, start=start_row):
            for c_index, value in enumerate(row_values, start=1):
                ws.cell(r_index, c_index).value = value
    for r_index in range(start_row + len(rows), old_max + 1):
        for c_index in range(1, 10):
            ws.cell(r_index, c_index).value = None

    auto_counts = {kind: 0 for kind in VALID_KR_TYPES}
    final_counts = {kind: 0 for kind in VALID_KR_TYPES}
    for row in rows:
        auto_counts[row[3]] += 1
        final_counts[row[4]] += 1

    summary['B6'] = auto_counts['병사']
    summary['B7'] = auto_counts['간부']
    summary['B8'] = auto_counts['미분류']
    summary['B9'] = len(rows)
    summary['C6'] = final_counts['병사']
    summary['C7'] = final_counts['간부']
    summary['C8'] = final_counts['미분류']
    summary['C9'] = len(rows)

    for row in summary.iter_rows(min_row=11, max_row=max(summary.max_row, 80), min_col=1, max_col=3):
        for cell in row:
            cell.value = None

    corrections = [row for row in rows if row[3] != row[4]]
    unknowns = [row for row in rows if row[4] == '미분류']
    cursor = 11
    summary.cell(cursor, 1).value = f'수정 필요 {len(corrections)}명'
    cursor += 1
    for row in corrections:
        summary.cell(cursor, 1).value = row[1]
        summary.cell(cursor, 2).value = f'{row[3]} → {row[4]}'
        summary.cell(cursor, 3).value = row[6]
        cursor += 1
    cursor += 1
    summary.cell(cursor, 1).value = f'미분류 유지 {len(unknowns)}명'
    cursor += 1
    for row in unknowns:
        summary.cell(cursor, 1).value = row[1]
        summary.cell(cursor, 2).value = row[6] or '신청 분야 추가 확인 필요'
        cursor += 1

    wb.save(workbook_path)
    type_map_path.write_text(type_map_text, encoding='utf-8')
    return True


def main() -> int:
    changed = sync_workbook()
    print('updated' if changed else 'no changes')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
